#!/usr/bin/env python3
"""Phases 3-5 acceptance — ingestion, cross-check, store, and the artefacts.

Runs offline against recorded payloads. What it proves:

  Phase 3  the marketplace pull parses, re-filters sponsored injections
           client-side, dedupes on source + listing ID across runs, and joins
           TrustMRR where matchable
  Phase 4  D6 and D7 are generated, and a re-run of D7 does not overwrite a
           human's stage, dates or notes
  Phase 5  D8 is generated for a promoted candidate and carries every flag with
           its evidence; outreach drafts are written to a review file

    python acquisition/tests/test_pipeline.py
"""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from acq.classify import HeuristicClassifier                     # noqa: E402
from acq.config import Config                                    # noqa: E402
from acq.fx import FixedRates                                    # noqa: E402
from acq.normalise import normalise                              # noqa: E402
from acq.rules import RuleEngine                                 # noqa: E402
from acq.store import Store                                      # noqa: E402
from acq.models import Candidate, NeglectCandidate               # noqa: E402
from acq.sources.fixture import FixtureSession                   # noqa: E402
from acq.sources.flippa import FlippaSource                      # noqa: E402
from acq.sources.trustmrr import TrustMrrRecord, cross_check     # noqa: E402
from acq.outputs import candidates_xlsx, diligence, outreach, tracker_xlsx  # noqa: E402

FIXTURE = Path(__file__).parent / "fixtures" / "flippa_page.json"
FX = FixedRates({"USD": 0.79, "GBP": 1.0})
GREEN, RED, RESET = "\033[32m", "\033[31m", "\033[0m"

results: list[tuple[bool, str, str]] = []


def expect(condition: bool, name: str, detail: str = "") -> None:
    results.append((bool(condition), name, detail))


def pull(cfg: Config) -> list[Candidate]:
    src = FlippaSource(cfg, session=FixtureSession.from_file(FIXTURE), fx=FX)
    src.delay = 0
    return [normalise(l, FX) for l in src.fetch()]


# ---------------------------------------------------------------- ingestion
def test_ingestion(cfg: Config) -> list[Candidate]:
    candidates = pull(cfg)
    ids = {c.source_id for c in candidates}

    expect("11500003" not in ids,
           "sponsored listing outside the price band is re-filtered client-side",
           "£379k ask survived the filter" if "11500003" in ids else "")
    expect(ids == {"11500001", "11500002", "11500004"},
           "the three in-band listings are kept", f"got {sorted(ids)}")

    shopify = next(c for c in candidates if c.source_id == "11500001")
    expect(abs(shopify.asking_price - 11500 * 0.79) < 0.01,
           "asking price converted to GBP at ingestion", f"£{shopify.asking_price:,.2f}")
    expect(shopify.original_currency == "USD" and shopify.fx_rate == 0.79,
           "original currency and rate recorded on the candidate")
    expect(abs(shopify.ttm_costs - (9600 - 7200) * 0.79) < 0.01,
           "costs derived from revenue minus profit", f"£{shopify.ttm_costs:,.2f}")
    expect(shopify.age_months and shopify.age_months > 24,
           "age derived from established_at", f"{shopify.age_months} months")
    return candidates


def test_dedupe(cfg: Config, tmp: Path) -> None:
    store = Store(tmp / "candidates.json", Candidate)
    first = store.upsert(pull(cfg))
    store.save()

    reloaded = Store(tmp / "candidates.json", Candidate)
    second = reloaded.upsert(pull(cfg))
    reloaded.save()

    expect(len(first) == 3, "first pull records three new candidates", f"{len(first)}")
    expect(second == [], "second pull of the same data adds nothing new", f"{second}")
    expect(len(reloaded) == 3, "store holds one record per source + listing ID",
           f"{len(reloaded)}")


def test_cross_check(cfg: Config) -> None:
    candidates = pull(cfg)
    records = [
        TrustMrrRecord(name="Shopify inventory alerts app", mrr=810,
                       url="https://inventoryalerts.app", currency="USD"),
        TrustMrrRecord(name="Agency time tracker", mrr=200,
                       url="https://agencytimetracker.com", currency="USD"),
    ]
    matched = cross_check(candidates, records, fx=FX)
    by_id = {c.source_id: c for c in candidates}

    expect(matched == 2, "both matchable candidates joined on product name", f"{matched}")
    expect(by_id["11500001"].trustmrr_match == "agrees",
           "a verified figure within tolerance agrees (800 vs 810)",
           by_id["11500001"].trustmrr_match)
    expect(by_id["11500004"].trustmrr_match == "disagrees",
           "a verified figure well below the listing disagrees (650 vs 200)",
           by_id["11500004"].trustmrr_match)
    expect(by_id["11500002"].trustmrr_match == "none",
           "an unmatched candidate is marked none, not silently verified")

    engine = RuleEngine(cfg, classifier=HeuristicClassifier())
    engine.screen_all(candidates)
    expect("UNVERIFIED" not in by_id["11500001"].flag_names(),
           "an agreeing TrustMRR match clears UNVERIFIED")
    expect("UNVERIFIED" in by_id["11500004"].flag_names(),
           "a disagreeing match does not clear UNVERIFIED")


def test_screening(cfg: Config) -> list[Candidate]:
    candidates = pull(cfg)
    RuleEngine(cfg, classifier=HeuristicClassifier()).screen_all(candidates)
    by_id = {c.source_id: c for c in candidates}

    expect(by_id["11500002"].disposition == "reject",
           "the ai-wrapper launched 5 months ago is rejected",
           by_id["11500002"].disposition)
    expect({"CATEGORY_EXCLUDED", "TOO_YOUNG"} <= set(by_id["11500002"].reject_reasons()),
           "rejected on both category and age",
           str(by_id["11500002"].reject_reasons()))
    expect("CHANNEL_RISK" in by_id["11500002"].flag_names(),
           "the viral-thread listing raises CHANNEL_RISK even though it is rejected")
    expect(by_id["11500001"].disposition == "review",
           "the organic Shopify app goes to review", by_id["11500001"].disposition)
    expect("CHANNEL_RISK" not in by_id["11500001"].flag_names(),
           "an app-store + organic channel does not raise CHANNEL_RISK")
    return candidates


# ------------------------------------------------------------------ outputs
def test_outputs(cfg: Config, tmp: Path) -> None:
    candidates = test_screening(cfg)
    review = [c for c in candidates if c.disposition == "review"]
    rejected = [c for c in candidates if c.disposition == "reject"]

    sheet = candidates_xlsx.write(review, tmp / "candidates.xlsx", rejected=rejected)
    wb = load_workbook(sheet)
    expect(wb.sheetnames == ["Shortlist", "Rejected"],
           "D6 has a shortlist and an audit tab for rejects", str(wb.sheetnames))
    ws = wb["Shortlist"]
    expect([c.value for c in ws[1]] == candidates_xlsx.HEADERS,
           "D6 columns match the brief")
    expect(ws.max_row == len(review) + 1, "one row per shortlisted candidate",
           f"{ws.max_row - 1} rows for {len(review)} candidates")
    expect(ws.freeze_panes == "A2" and ws.auto_filter.ref, "header frozen and filterable")
    scores = [ws.cell(row=r, column=candidates_xlsx.HEADERS.index("Score") + 1).value
              for r in range(2, ws.max_row + 1)]
    expect(scores == sorted(scores, reverse=True), "rows are ranked by score descending")
    flagged = [r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=14).value]
    expect(all(ws.cell(row=r, column=1).fill.fgColor.rgb.endswith(("FDE2E1", "DFF3E3"))
               for r in flagged),
           "every flagged or verified row is colour-coded")

    tracker = tracker_xlsx.write(review, tmp / "deal-tracker.xlsx")
    tw = load_workbook(tracker)
    expect("Pipeline" in tw.sheetnames and "Stages" in tw.sheetnames,
           "D7 has a pipeline sheet and a stage legend", str(tw.sheetnames))
    tws = tw["Pipeline"]
    expect(len(tws.data_validations.dataValidation) == 1,
           "stage column is a dropdown, so the pipeline can't drift into free text")

    # A human edits the tracker; the next run must not clobber it.
    stage_col = tracker_xlsx.HEADERS.index("Stage") + 1
    notes_col = tracker_xlsx.HEADERS.index("Notes") + 1
    action_col = tracker_xlsx.HEADERS.index("Next action") + 1
    tws.cell(row=2, column=stage_col, value="responded")
    tws.cell(row=2, column=notes_col, value="Seller sent Stripe screenshots — asked for export")
    tws.cell(row=2, column=action_col, value="chase processor export")
    edited_id = tws.cell(row=2, column=1).value
    tw.save(tracker)

    tracker_xlsx.write(review, tracker)
    re_read = load_workbook(tracker)["Pipeline"]
    row = next(r for r in range(2, re_read.max_row + 1)
               if re_read.cell(row=r, column=1).value == edited_id)
    expect(re_read.cell(row=row, column=stage_col).value == "responded",
           "D7 re-run preserves a human-set stage")
    expect(re_read.cell(row=row, column=notes_col).value.startswith("Seller sent"),
           "D7 re-run preserves human notes")
    expect(re_read.cell(row=row, column=action_col).value == "chase processor export",
           "D7 re-run preserves the next action")


def test_diligence_and_outreach(cfg: Config, tmp: Path) -> None:
    candidates = test_screening(cfg)
    promoted = next(c for c in candidates if c.disposition == "review")
    promoted.status = "contacted"

    path = diligence.write(promoted, tmp / "diligence")
    text = path.read_text()
    expect(path.exists(), "D8 pack written for a promoted candidate")
    for section in ("Computed metrics", "Flags", "Triage", "Data request", "Reconciliation"):
        expect(section in text, f"D8 contains the '{section}' section")
    for f in promoted.flags:
        expect(f["flag"] in text, f"D8 lists the {f['flag']} flag")
        if f["evidence"]:
            expect(f["evidence"][:40] in text, f"D8 shows the evidence for {f['flag']}")
    expect("Where did the customers come from?" in text,
           "D8 carries the four triage questions from the playbook")

    neglect = [NeglectCandidate(slug="simple-booking-calendar", name="Simple Booking Calendar",
                                url="https://wordpress.org/plugins/simple-booking-calendar/",
                                author="Helen Griffiths", active_installs=40000,
                                months_since_update=43.2, rating=4.4, num_ratings=214,
                                author_profile="https://profiles.wordpress.org/hgriffiths/",
                                support_threads=18, support_threads_resolved=2,
                                neglect_score=8.9)]
    drafts = outreach.write(neglect, [promoted], tmp / "outreach-drafts.md", sender="John")
    body = drafts.read_text()
    expect("40,000 active installs" in body, "off-market draft names the install count")
    expect("Helen" in body, "off-market draft addresses the author by first name")
    expect("I don't need a spreadsheet" in body,
           "marketplace draft pre-empts the seller-built spreadsheet")
    expect("valuation" not in body.lower().split("do not attach a valuation")[-1],
           "no valuation appears in any draft body")
    expect("has no send path" in body, "the drafts file states that nothing was sent")


def main() -> int:
    cfg = Config.load()
    tmp = Path(tempfile.mkdtemp(prefix="acq-test-"))
    print("INGESTION → SCREENING → OUTPUTS — PHASE 3-5 ACCEPTANCE")
    print("=" * 72)
    try:
        test_ingestion(cfg)
        test_dedupe(cfg, tmp)
        test_cross_check(cfg)
        test_outputs(cfg, tmp)
        test_diligence_and_outreach(cfg, tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    failed = 0
    print()
    for ok, name, detail in results:
        mark = f"{GREEN}PASS{RESET}" if ok else f"{RED}FAIL{RESET}"
        suffix = f"  ({detail})" if detail and not ok else ""
        print(f"  [{mark}] {name}{suffix}")
        failed += 0 if ok else 1

    print("=" * 72)
    if failed:
        print(f"{RED}FAILED{RESET} — {failed} of {len(results)} checks")
        return 1
    print(f"{GREEN}PASSED{RESET} — {len(results)} checks")
    return 0


if __name__ == "__main__":
    sys.exit(main())
