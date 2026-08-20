"""D6 — the ranked candidate sheet. One row per candidate.

Everything the engine computed, with the flags visible. The engine sorts; the
human decides. Rejected candidates are written to a second tab rather than
discarded, because "why did we never see this one?" is a question the sheet
should be able to answer a month later.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..models import Candidate
from .style import (BORDER, FLAG_FILL, MONEY, MONEY_P, NUMBER, RATIO,
                    REJECT_FILL, VERIFIED_FILL, autosize, write_header)

HEADERS = [
    "Source", "URL", "Name", "Category", "Asking price", "MRR", "TTM revenue",
    "ARPU", "Price/revenue", "Price/profit", "Payback months", "Trend ratio",
    "Verification", "Flags", "Score", "First seen", "Status",
]
WIDTHS = {"URL": 34, "Name": 30, "Flags": 46, "Category": 16, "Verification": 16}
FORMATS = {
    "Asking price": MONEY, "MRR": MONEY, "TTM revenue": MONEY, "ARPU": MONEY_P,
    "Price/revenue": RATIO, "Price/profit": RATIO, "Payback months": NUMBER,
    "Trend ratio": "0.00", "Score": "0.0",
}


def verification_label(c: Candidate) -> str:
    if c.has_verified_revenue and c.trustmrr_match == "agrees":
        return "verified + TrustMRR"
    if c.trustmrr_match == "agrees":
        return "TrustMRR agrees"
    if c.trustmrr_match == "disagrees":
        return "CONTRADICTED by TrustMRR"
    if c.has_verified_revenue:
        return "listing verified"
    return "unverified"


def row_for(c: Candidate) -> list:
    return [
        c.source, c.url, c.name, c.category, c.asking_price, c.mrr, c.ttm_revenue,
        c.arpu, c.price_to_revenue, c.price_to_profit, c.payback_months, c.trend_ratio,
        verification_label(c), ", ".join(c.flag_names()), c.score, c.first_seen, c.status,
    ]


def _write_sheet(ws, candidates: list[Candidate], *, rejected: bool) -> None:
    write_header(ws, HEADERS)
    for r, c in enumerate(candidates, start=2):
        values = row_for(c)
        if rejected:
            values[13] = "; ".join(c.reject_reasons()) or values[13]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = BORDER
            title = HEADERS[col - 1]
            if title in FORMATS:
                cell.number_format = FORMATS[title]
            if title in ("URL", "Flags", "Name"):
                cell.alignment = Alignment(wrap_text=False, vertical="center")

        # Conditional formatting, per the brief: red for any flag, green for
        # verified revenue. Verification wins the tie — a verified figure is the
        # one thing that changes what the flags are worth.
        verified = c.has_verified_revenue or c.trustmrr_match == "agrees"
        fill = None
        if rejected:
            fill = REJECT_FILL
        elif verified:
            fill = VERIFIED_FILL
        elif c.flags:
            fill = FLAG_FILL
        if fill:
            for col in range(1, len(HEADERS) + 1):
                ws.cell(row=r, column=col).fill = fill

        if c.url:
            link = ws.cell(row=r, column=2)
            link.hyperlink = c.url
            link.font = Font(color="0B62A4", underline="single")

    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(ws.max_row, 1)}"
    autosize(ws, HEADERS, WIDTHS)


def write(candidates: list[Candidate], path: str | Path,
          rejected: list[Candidate] | None = None) -> Path:
    ranked = sorted(candidates, key=lambda c: c.score, reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Shortlist"
    _write_sheet(ws, ranked, rejected=False)

    if rejected:
        rs = wb.create_sheet("Rejected")
        _write_sheet(rs, sorted(rejected, key=lambda c: c.score, reverse=True), rejected=True)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
