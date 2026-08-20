"""Shared spreadsheet styling. One place, so D6 and D7 look like one system."""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F2933")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
FLAG_FILL = PatternFill("solid", fgColor="FDE2E1")       # any flag raised
VERIFIED_FILL = PatternFill("solid", fgColor="DFF3E3")   # processor-verified revenue
REJECT_FILL = PatternFill("solid", fgColor="EFEFEF")     # hard-rejected, kept for audit
THIN = Side(style="thin", color="D8DEE4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY = '£#,##0'
MONEY_P = '£#,##0.00'
RATIO = '0.00"x"'
NUMBER = '#,##0.0'


def write_header(ws: Worksheet, headers: list[str], row: int = 1) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.row_dimensions[row].height = 30


def autosize(ws: Worksheet, headers: list[str], widths: dict[str, int] | None = None,
             default: int = 14, cap: int = 52) -> None:
    widths = widths or {}
    for idx, title in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        if title in widths:
            ws.column_dimensions[letter].width = widths[title]
            continue
        longest = max((len(str(ws.cell(row=r, column=idx).value or ""))
                       for r in range(1, ws.max_row + 1)), default=0)
        ws.column_dimensions[letter].width = min(max(default, longest + 2), cap)
